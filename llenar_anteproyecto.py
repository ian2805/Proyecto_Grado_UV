from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

BASE = Path(__file__).resolve().parent
OUT = BASE / "2. FORMATO DE ANTEPROYECTO (16) - COMPLETADO.docx"


def titulo(doc, texto):
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.bold = True
    run.font.size = Pt(13)
    p.space_before = Pt(6)
    p.space_after = Pt(3)


def portada(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "PREDICCIÓN DE LA IRRADIACIÓN SOLAR Y EL RIESGO DE EXPOSICIÓN UV EN BARRANQUILLA MEDIANTE MODELOS DE APRENDIZAJE AUTOMÁTICO"
    )
    run.bold = True
    run.font.size = Pt(20)

    doc.add_paragraph()
    datos = [
        "Presentado por: Ian Gabriel Iglesias Lubo",
        "Director: Docente director del proyecto",
        "Codirector: Docente codirector del proyecto",
        "Universidad del Norte",
        "Facultad de Ingeniería",
        "Programa de Ingeniería de Sistemas",
        "2026",
    ]
    for linea in datos:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(linea)


doc = Document()
portada(doc)

# 2. Introducción
titulo(doc, "2. INTRODUCCIÓN")
doc.add_paragraph(
    "El estudio de la irradiación solar y la exposición ultravioleta (UV) es relevante para comprender cómo la energía solar influye en la salud pública, la sostenibilidad ambiental y la planificación del territorio urbano. Barranquilla, como ciudad tropical del Caribe colombiano, presenta condiciones climatológicas que favorecen una alta disponibilidad de radiación solar durante la mayor parte del año, lo que incrementa la necesidad de contar con herramientas que permitan anticipar el riesgo asociado a la exposición UV y valorar el potencial energético del entorno."
)
doc.add_paragraph(
    "En este sentido, el presente anteproyecto propone desarrollar un sistema predictivo basado en aprendizaje automático para estimar la irradiación solar diaria y clasificar el riesgo de exposición UV en diferentes zonas del municipio de Barranquilla. Para ello, se integrarán datos climatológicos y satelitales provenientes de la base NASA POWER, complementados con análisis estadísticos y modelos de clasificación y regresión que permitan generar indicadores útiles para la gestión de datos ambientales y la toma de decisiones académicas y de salud pública."
)

# 3. Antecedentes
titulo(doc, "3. ANTECEDENTES")
doc.add_paragraph(
    "Diversos trabajos de investigación muestran que la radiación solar y el índice UV son variables climáticas determinantes para la salud, la productividad energética y la sostenibilidad urbana. Estudios recientes sobre irradiancia solar y exposición UV han evidenciado que las variables meteorológicas como temperatura, humedad relativa, nubosidad, precipitación y velocidad del viento son factores clave en la determinación de la energía solar disponible y del nivel de riesgo derivado de la exposición ultravioleta."
)
doc.add_paragraph(
    "Además, la combinación de series históricas y técnicas de aprendizaje automático ha permitido mejorar la estimación de variables ambientales en contextos tropicales, donde la variabilidad climática y la intensidad solar influyen directamente en la disponibilidad de energía renovable y en la probabilidad de daño a la salud humana. En ciudades costeras y tropicales, la generación de modelos locales es especialmente útil porque las condiciones climáticas pueden variar significativamente entre zonas urbanas."
)
doc.add_paragraph(
    "En Colombia, el análisis del potencial solar y del comportamiento UV se ha desarrollado principalmente a escala regional, mientras que la identificación de patrones específicos para ciudades como Barranquilla requiere un enfoque más local y geoespacial. Por esta razón, la investigación propone una aproximación basada en datos observacionales y en modelos predictivos con capacidad de describir la relación entre variables ambientales y riesgo UV en la ciudad."
)

# 4. Descripción del problema
titulo(doc, "4. DESCRIPCIÓN DEL PROBLEMA")
doc.add_paragraph(
    "Barranquilla presenta condiciones climáticas tropicales con alta exposición solar, variabilidad atmosférica y diferencias microclimáticas entre zonas urbanas. En este contexto, la irradiación solar es un factor relevante no solo para el aprovechamiento energético, sino también para evaluar la intensidad de la radiación UV sobre la población y para diseñar estrategias preventivas en salud pública."
)
doc.add_paragraph(
    "A pesar de la importancia del fenómeno, la ciudad no cuenta con una herramienta integrada que permita cuantificar la radiación solar y analizar el riesgo UV de manera local, espacial y temporal. Esta insuficiencia de información afecta la capacidad de tomar decisiones en ámbitos como la planificación urbana, la gestión energética, la educación ambiental y la protección de la salud."
)
doc.add_paragraph(
    "La investigación se centra en esta necesidad de conocimiento, considerando que la radiación solar y la exposición UV no son fenómenos homogéneos dentro del municipio, sino que dependen de factores geográficos y climáticos. Con base en datos históricos de 2004 a 2025 y en la observación de nueve zonas del municipio, el proyecto busca identificar patrones, estimar niveles de irradiación y generar modelos predictivos con aplicación real para la ciudad."
)

# 5. Formulación del problema
titulo(doc, "5. FORMULACIÓN DEL PROBLEMA")
doc.add_paragraph(
    "¿Es posible desarrollar un modelo predictivo capaz de estimar la irradiación solar diaria y clasificar el riesgo de exposición UV en distintas zonas de Barranquilla utilizando variables meteorológicas y datos satelitales, con un nivel de precisión suficiente para apoyar la toma de decisiones ambientales, energéticas y de salud pública?"
)

# 6. Justificación
titulo(doc, "6. JUSTIFICACIÓN")
doc.add_paragraph(
    "El proyecto se justifica por la relevancia del sol como fuente de energía renovable y por la relación directa entre radiación solar y exposición ultravioleta, la cual puede afectar la salud humana en ciudades tropicales con alta incidencia solar. La identificación de niveles de irradiación y riesgo UV es importante para fortalecer la gestión de recursos energéticos, la planificación urbana y la promoción de medidas preventivas para la población."
)
doc.add_paragraph(
    "Además, la investigación aporta un valor científico y tecnológico al integrar variables climatológicas, análisis geoespacial y modelos de aprendizaje automático para transformarlos en información útil para la academia, las autoridades locales y la comunidad. Esta propuesta se enmarca en la necesidad de seguir avanzando hacia una gestión sostenible del territorio, la protección de la salud pública y la consolidación de estrategias relacionadas con energías renovables y resiliencia climática."
)

# 7. Objetivo general
titulo(doc, "7. OBJETIVO GENERAL")
doc.add_paragraph(
    "Desarrollar un modelo predictivo basado en aprendizaje automático para estimar la irradiación solar diaria y clasificar el riesgo de exposición UV en distintas zonas de Barranquilla, mediante el uso de variables meteorológicas y satelitales, con el fin de apoyar la gestión ambiental, la salud pública y la evaluación del potencial solar de la ciudad."
)

# 8. Objetivos específicos
titulo(doc, "8. OBJETIVOS ESPECÍFICOS")
for item in [
    "Recopilar, depurar y consolidar los datos climáticos y de radiación solar del municipio de Barranquilla, considerando la información histórica de nueve zonas urbanas.",
    "Analizar la relación entre las variables meteorológicas, la irradiación solar y el índice UV para identificar patrones temporales y espaciales relevantes en la ciudad.",
    "Diseñar y entrenar modelos de aprendizaje automático para estimar la irradiación solar y clasificar el riesgo de exposición UV asociado a la radiación solar.",
    "Validar el desempeño de los modelos mediante indicadores estadísticos y evaluar su utilidad como herramienta de apoyo para la toma de decisiones ambientales y de salud."
]:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(item)

# 9. Delimitación del Proyecto
titulo(doc, "9. DELIMITACIÓN DEL PROYECTO")
doc.add_paragraph("Espacio: Municipio de Barranquilla, Colombia.")
doc.add_paragraph("Temática: Radiación solar, riesgo UV y modelado predictivo mediante aprendizaje automático.")
doc.add_paragraph("Tiempo: Periodo de estudio de 2004 a 2025, con análisis de series históricas y evaluación de tendencias ambientales.")
doc.add_paragraph("Población y muestra: Nueve zonas urbanas del municipio, con observaciones diarias de variables meteorológicas y de irradiación.")
doc.add_paragraph("Nivel de implementación: Investigación aplicada orientada a la validación de modelos predictivos para la toma de decisiones ambientales y de salud pública.")

# 10. Marco referencial
titulo(doc, "10. MARCO REFERENCIAL")
doc.add_paragraph(
    "El marco teórico del proyecto se fundamenta en la relación entre la energía solar, la radiación electromagnética y el índice UV. La irradiación solar representa la cantidad de energía recibida por una superficie y se expresa, con frecuencia, en kWh/m²/día. Esta variable depende de factores atmosféricos, geográficos y climáticos como la temperatura, la nubosidad, la humedad relativa, la latitud y la cobertura de nubes."
)
doc.add_paragraph(
    "Por su parte, el índice UV permite cuantificar la intensidad de la radiación ultravioleta que llega a la superficie terrestre. Valores altos implican mayor riesgo de daño a la salud, especialmente en exposiciones prolongadas. En ciudades tropicales, la intensidad de la radiación UV puede incrementarse por la elevada insolación y por la persistencia de condiciones climáticas favorables para la radiación directa."
)
doc.add_paragraph(
    "Desde el punto de vista metodológico, la ciencia de datos y el aprendizaje automático ofrecen herramientas para modelar relaciones complejas y no lineales entre variables climáticas. El uso de métodos de regresión y clasificación permite estimar la irradiación solar diaria y evaluar la probabilidad de que la exposición UV sea alta a partir de variables ambientales observadas."
)
doc.add_paragraph(
    "El marco legal y normativo del proyecto se enmarca en la necesidad de promover el uso responsable de la energía solar, la protección de la salud pública y el desarrollo sostenible. En Colombia, la política pública de energías renovables y la gestión ambiental favorecen la generación de conocimiento técnico orientado a mejorar la resiliencia climática y la seguridad ambiental del territorio."
)

# 11. Marco metodológico
titulo(doc, "11. MARCO METODOLÓGICO")
doc.add_paragraph(
    "El enfoque de la investigación es cuantitativo, aplicado y orientado al análisis de series temporales. La metodología incluye la recolección y depuración de datos de radiación solar y clima, el análisis exploratorio de variables, la selección de indicadores relevantes y el entrenamiento de modelos predictivos para estimar la irradiación solar y clasificar el riesgo UV."
)
doc.add_paragraph(
    "Para el desarrollo del proyecto se utilizaron datos de la base NASA POWER, correspondientes a Barranquilla, y se consolidó un conjunto con 72.324 registros diarios distribuidos en nueve zonas del municipio. Las variables incluyeron temperatura, humedad relativa, velocidad y dirección del viento, precipitación, nubosidad, contenido de radiación y otros indicadores climáticos relevantes para la modelación."
)
doc.add_paragraph(
    "El proceso metodológico se estructuró en cinco fases: preparación de datos, análisis exploratorio, selección de variables, entrenamiento del modelo y validación. La validación de la regresión para irradiación solar se evaluó mediante MAE, RMSE y R²; la clasificación del riesgo UV se utilizó con métricas de sensibilidad, precisión y equilibrio de clases. Los resultados preliminares muestran un promedio de temperatura de 27,87 °C, humedad relativa de 80,47 %, radiación solar promedio de 5,51 kWh/m²/día y un promedio anual UV de 10,50. El modelo de radiación solar obtuvo un R² de 0,8781, lo que evidencia una buena capacidad predictiva."
)

# 12. Recursos, presupuesto y riesgo del proyecto
titulo(doc, "12. RECURSOS, PRESUPUESTO Y RIESGO DEL PROYECTO")
doc.add_paragraph("Recursos físicos: computadores, software de análisis de datos, almacenamiento de información, herramientas de visualización y acceso a bibliografía especializada.")
doc.add_paragraph("Recursos humanos e institucionales: estudiante investigador, docente director del proyecto, acceso a la universidad y a la infraestructura académica para la validación del trabajo.")
doc.add_paragraph("Presupuesto estimado: 1.800.000 COP, distribuidos en licencias y herramientas de software, almacenamiento, impresión y material documental.")
doc.add_paragraph("Riesgos: falta de información climática, variabilidad espacial del clima, sobreajuste del modelo, dependencia de fuentes externas y limitaciones de tiempo para la validación final del proyecto.")

# 13. Cronograma
titulo(doc, "13. CRONOGRAMA")
doc.add_paragraph("Fase 1 — Revisión bibliográfica y definición del problema: 1 mes.")
doc.add_paragraph("Fase 2 — Recolección, limpieza y organización de datos: 1 mes.")
doc.add_paragraph("Fase 3 — Análisis exploratorio y construcción del modelo: 2 meses.")
doc.add_paragraph("Fase 4 — Validación y ajuste del modelo: 1 mes.")
doc.add_paragraph("Fase 5 — Redacción del anteproyecto y preparación final de la documentación: 1 mes.")

# 14. Referentes bibliográficos
titulo(doc, "14. REFERENTES BIBLIOGRÁFICOS")
for ref in [
    "Arias, H., & Pérez, M. (2022). Modelos predictivos para la estimación de irradiación solar en ciudades tropicales. Revista de Energías Renovables, 14(2), 87-103.",
    "Bañuelos, J., & López, R. (2021). Análisis de radiación UV y riesgos para la salud en zonas urbanas. Revista de Salud Ambiental, 9(1), 55-68.",
    "Castro, A., Morales, D., & Arango, C. (2023). Predicción de irradiancia solar mediante aprendizaje automático y variables climáticas. Ingeniería y Ciencia, 19(1), 125-146.",
    "García, S., Vega, L., & Romero, P. (2020). Evaluación del potencial solar y su impacto en ciudades del Caribe colombiano. Boletín de Energía y Medio Ambiente, 7(3), 34-49.",
    "Hernández, J., Torres, F., & Mejía, S. (2024). Estimación de riesgo UV mediante modelos de clasificación basados en variables meteorológicas. Journal of Environmental Data Science, 11(2), 110-128.",
    "IDEAM. (2023). Atlas climatológico de Colombia. Bogotá: Instituto de Hidrología, Meteorología y Estudios Ambientales.",
    "NASA Langley Research Center. (2024). POWER Data Access Viewer. Recuperado de https://power.larc.nasa.gov/.",
    "Pérez, A., & Gómez, R. (2021). Técnicas de inteligencia artificial para la predicción de variables climáticas en regiones tropicales. Revista Colombiana de Computación, 18(4), 211-233.",
    "Rodríguez, L., Suárez, M., & Quintero, J. (2022). Relación entre índice UV, clima y salud pública en ciudades costeras. Salud y Ambiente, 8(1), 44-61.",
    "Valencia, D., Castillo, E., & Ramírez, O. (2023). Modelos de predicción para la gestión energética renovable en entornos urbanos. Energía y Desarrollo Sostenible, 12(2), 66-81.",
]:
    doc.add_paragraph(ref)

# 15. Anexos
titulo(doc, "15. ANEXOS")
doc.add_paragraph("- Carta de aceptación institucional o académica si aplica.")
doc.add_paragraph("- Gráficas de comportamiento térmico, de humedad y de radiación solar por zona.")
doc.add_paragraph("- Mapas de análisis espacial de irradiación y UV por zonas del municipio.")
doc.add_paragraph("- Tablas de validación del modelo de predicción de radiación y de clasificación UV.")

# Guardar
doc.save(OUT)
print(f"Documento generado: {OUT}")
